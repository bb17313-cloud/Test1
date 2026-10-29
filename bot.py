# -*- coding: utf-8 -*-
"""
بوت تلغرام (يعمل على GitHub Actions) - مصادر مجانية بدون أي API مدفوع
------------------------------------------------------------------
1) ترجمة أي نص / ملف / صورة / رابط إلى العربية   (Google Translate المجاني)
2) تلخيص أي مستند (PDF, Word, PowerPoint, Excel, TXT, HTML, صور)  (تلخيص استخراجي محلي)
3) تحليل أي سهم: النشاط، نتائج الربع، المستهدفات، المشاريع والعقود، الاتجاه،
   زيادة/خفض رأس المال، التوزيعات، التقسيم العكسي، الخسائر المتراكمة، الأخبار
   (Yahoo Finance + SEC EDGAR - بدون مفاتيح)

الأسرار المطلوبة: BOT_TOKEN, CHAT_ID
اختياري: SEC_USER_AGENT (مثال: "MyBot myemail@gmail.com")
"""
import os
import re
import sys
import time
import math
import logging
import tempfile
import xml.etree.ElementTree as ET
from collections import Counter
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime

import requests
from bs4 import BeautifulSoup

# ──────────────────────────── الإعدادات ────────────────────────────
BOT_TOKEN = os.environ["BOT_TOKEN"].strip()
CHAT_ID = str(os.environ["CHAT_ID"]).strip()
RUN_SECONDS = int(os.environ.get("RUN_SECONDS") or 21000)
SEC_UA = os.environ.get("SEC_USER_AGENT") or "TelegramStockBot admin@example.com"

API = f"https://api.telegram.org/bot{BOT_TOKEN}"
FILE_API = f"https://api.telegram.org/file/bot{BOT_TOKEN}"
WEB_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                  "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Accept-Language": "en-US,en;q=0.9,ar;q=0.8",
}
SEC_HEADERS = {"User-Agent": SEC_UA, "Accept-Encoding": "gzip, deflate"}

logging.basicConfig(level=logging.INFO, format="%(asctime)s | %(levelname)s | %(message)s")
log = logging.getLogger("bot")

HELP = (
    "👋 أهلاً! أرسل لي:\n\n"
    "• أي نص ← أترجمه إلى العربية\n"
    "• ملف (PDF / Word / PowerPoint / Excel / TXT / HTML) ← ألخّصه بالعربية\n"
    "   ↳ اكتب «ترجم» في وصف الملف لترجمته كاملاً\n"
    "• رابط مقال ← ألخّصه (أو «ترجم + الرابط» لترجمته)\n"
    "• صورة ← أستخرج النص وأترجمه (أو «لخص» في الوصف)\n"
    "• رمز سهم مثل AAPL أو رابط TradingView أو «تحليل AAPL» ← تحليل شامل\n\n"
    "الأوامر: /tr ترجمة  |  /sum تلخيص  |  /stock تحليل سهم\n"
    "يمكنك أيضاً الرد على أي رسالة أو ملف بكلمة «ترجم» أو «لخص»."
)


# ──────────────────────────── تلغرام ────────────────────────────
def tg(method, files=None, timeout=60, **params):
    params = {k: v for k, v in params.items() if v is not None}
    for attempt in range(3):
        try:
            r = requests.post(f"{API}/{method}", data=params, files=files, timeout=timeout)
            j = r.json()
            if j.get("ok"):
                return j["result"]
            log.warning("tg %s failed: %s", method, j)
            if r.status_code == 429:
                time.sleep(int(j.get("parameters", {}).get("retry_after", 3)) + 1)
                continue
            return None
        except Exception as e:  # noqa
            log.warning("tg %s error: %s", method, e)
            time.sleep(2)
    return None


def send(chat_id, text, reply_to=None):
    text = (text or "").strip()
    if not text:
        return
    while text:
        if len(text) <= 3900:
            part, text = text, ""
        else:
            cut = text.rfind("\n", 0, 3900)
            cut = cut if cut > 1500 else 3900
            part, text = text[:cut], text[cut:].lstrip("\n")
        tg("sendMessage", chat_id=chat_id, text=part, reply_to_message_id=reply_to,
           disable_web_page_preview="true")
        reply_to = None


def send_doc(chat_id, content, filename, caption=None):
    data = content.encode("utf-8") if isinstance(content, str) else content
    tg("sendDocument", files={"document": (filename, data)}, chat_id=chat_id,
       caption=caption, timeout=120)


def typing(chat_id):
    tg("sendChatAction", chat_id=chat_id, action="typing", timeout=15)


def download_file(file_id):
    info = tg("getFile", file_id=file_id)
    if not info or not info.get("file_path"):
        raise RuntimeError("تعذّر تنزيل الملف (حدّ تلغرام للبوتات 20MB)")
    r = requests.get(f"{FILE_API}/{info['file_path']}", timeout=120)
    r.raise_for_status()
    return r.content


# ──────────────────────────── أدوات النص ────────────────────────────
def is_arabic(text, threshold=0.5):
    letters = re.findall(r"[^\W\d_]", text or "")
    if not letters:
        return False
    ar = sum(1 for c in letters if "\u0600" <= c <= "\u06FF")
    return ar / len(letters) >= threshold


def chunk_text(text, limit=4000):
    chunks, cur = [], ""
    for line in text.split("\n"):
        while len(line) > limit:
            cut = max(line.rfind(". ", 0, limit), line.rfind(" ", 0, limit))
            cut = cut if cut > limit // 2 else limit
            if cur:
                chunks.append(cur)
                cur = ""
            chunks.append(line[:cut])
            line = line[cut:].lstrip()
        if len(cur) + len(line) + 1 > limit and cur:
            chunks.append(cur)
            cur = ""
        cur = f"{cur}\n{line}" if cur else line
    if cur:
        chunks.append(cur)
    return chunks


# ── حد معدل طلبات الترجمة (Google يسمح ≈ 5/ثانية)
_last_gtx = 0.0
_GTX_MIN_GAP = 0.35  # ثانية بين كل طلب


def _gtx_wait():
    global _last_gtx
    gap = time.time() - _last_gtx
    if gap < _GTX_MIN_GAP:
        time.sleep(_GTX_MIN_GAP - gap)
    _last_gtx = time.time()


def _is_rate_limit(err):
    s = str(err).lower()
    return any(k in s for k in ("too many requests", "429", "rate", "quota", "5 requests per second"))


def _gtx(chunk, sl):
    _gtx_wait()
    r = requests.post(
        "https://translate.googleapis.com/translate_a/single",
        params={"client": "gtx", "sl": sl, "tl": "ar", "dt": "t"},
        data={"q": chunk}, headers=WEB_HEADERS, timeout=30)
    if r.status_code == 429:
        raise RuntimeError(f"429 rate limit: {r.text[:200]}")
    r.raise_for_status()
    data = r.json()
    return "".join(x[0] for x in data[0] if x and x[0])


def _translate_chunk(chunk, soft=False):
    """soft=True: عند الفشل يرجع النص الأصلي بدل رفع خطأ."""
    if not chunk.strip():
        return chunk
    needs = bool(re.search(r"[^\W\d_]", chunk)) and not is_arabic(chunk, 0.9)
    has_latin = bool(re.search(r"[A-Za-z]{3}", chunk))
    best, last_err = None, None

    # محاولة واحدة بـ auto أولاً؛ en فقط إذا فشلت أو النتيجة غير عربية
    for sl in ("auto", "en"):
        out = None
        for attempt in range(4):
            try:
                out = _gtx(chunk, sl)
                break
            except Exception as e:  # noqa
                last_err = e
                log.warning("gtx %s attempt %s: %s", sl, attempt + 1, e)
                if _is_rate_limit(e):
                    time.sleep(3 * (attempt + 1))  # 3، 6، 9، 12 ثانية
                else:
                    time.sleep(1.2 * (attempt + 1))
        if out and out.strip() and (out.strip() != chunk.strip() or not needs):
            best = out
            if is_arabic(out, 0.6) or not has_latin:
                return out
        if best:
            break  # لا نجرب en إذا حصلنا نتيجة مقبولة من auto

    if best:
        return best

    # احتياطي: deep_translator
    try:
        from deep_translator import GoogleTranslator
        time.sleep(1.5)
        for sl in ("auto", "en"):
            out = GoogleTranslator(source=sl, target="ar").translate(chunk)
            if out and out.strip() != chunk.strip():
                return out
    except Exception as e:  # noqa
        last_err = e
        log.warning("deep_translator: %s", e)

    if soft:
        log.warning("translate soft-fail, returning original. last_err=%s", last_err)
        return chunk
    raise RuntimeError(
        f"فشلت الترجمة مؤقتاً (حد طلبات Google). انتظر دقيقة ثم أعد المحاولة. "
        f"{str(last_err or '')[:120]}"
    )


def translate_ar(text):
    text = (text or "").strip()
    if not text or is_arabic(text, 0.9):
        return text
    parts = []
    chunks = chunk_text(text, 1500)  # قطع أصغر = أقل ضغط
    for i, c in enumerate(chunks):
        if i > 0:
            time.sleep(0.4)
        parts.append(_translate_chunk(c, soft=True))
    return "\n".join(parts)


def translate_lines(lines):
    """ترجمة قائمة أسطر مع الحفاظ على الترتيب."""
    lines = [l.strip() for l in lines]
    result, group, idx = [None] * len(lines), [], []

    def flush():
        if not group:
            return
        joined = "\n".join(group)
        out = _translate_chunk(joined, soft=True).split("\n")
        if len(out) == len(group):
            for i, o in zip(idx, out):
                result[i] = o.strip()
        else:
            for i, g in zip(idx, group):
                time.sleep(0.3)
                result[i] = _translate_chunk(g, soft=True).strip()
        group.clear()
        idx.clear()

    size = 0
    for i, l in enumerate(lines):
        if not l or is_arabic(l, 0.9):
            result[i] = l
            continue
        if size + len(l) > 2800:
            flush()
            size = 0
            time.sleep(0.5)
        group.append(l.replace("\n", " "))
        idx.append(i)
        size += len(l) + 1
    flush()
    return [r if r is not None else l for r, l in zip(result, lines)]


# ──────────────────────────── التلخيص ────────────────────────────
def split_sentences(text):
    text = re.sub(r"[ \t\u00a0]+", " ", text or "")
    text = re.sub(r"(?<![.!?؟:\n])\n(?!\n)", " ", text)  # دمج أسطر PDF المكسورة
    parts = re.split(r"(?<=[.!?؟])\s+|\n+", text)
    return [p.strip() for p in parts if 30 <= len(p.strip()) <= 700]


def summarize_extractive(text, n=None):
    sents = split_sentences(text)
    if len(sents) <= 6:
        return sents
    n = n or max(5, min(14, len(sents) // 10))
    toks = [set(re.findall(r"\w{3,}", s.lower())) for s in sents]
    df = Counter(w for t in toks for w in t)
    limit = max(3, int(len(sents) * 0.4))
    freq = Counter()
    for s in sents:
        for w in re.findall(r"\w{3,}", s.lower()):
            if df[w] <= limit:  # استبعاد الكلمات الشائعة جداً
                freq[w] += 1
    scores = []
    for i, s in enumerate(sents):
        words = re.findall(r"\w{3,}", s.lower())
        if not words:
            scores.append(0)
            continue
        sc = sum(freq[w] for w in set(words)) / (len(words) ** 0.7 + 1)
        if i < 3:
            sc *= 1.3
        if re.search(r"\d", s):
            sc *= 1.1
        scores.append(sc)
    chosen = []
    for i in sorted(range(len(sents)), key=lambda k: scores[k], reverse=True):
        if len(chosen) >= n:
            break
        if any(len(toks[i] & toks[j]) / max(1, len(toks[i] | toks[j])) > 0.55 for j in chosen):
            continue
        chosen.append(i)
    return [sents[i] for i in sorted(chosen)]


def key_figures(text, exclude, limit=5):
    pat = re.compile(r"(\$|€|£|%|\d[\d,\.]*\s?(million|billion|مليون|مليار|M|B)\b)", re.I)
    out = []
    for s in split_sentences(text):
        if s not in exclude and pat.search(s):
            out.append(s)
        if len(out) >= limit:
            break
    return out


def build_summary(text, title):
    words = len(text.split())
    sents = summarize_extractive(text)
    if not sents:
        return "⚠️ النص قصير أو غير قابل للتلخيص."
    ar = translate_lines(sents)
    out = [f"📝 ملخص: {title}", f"(الحجم الأصلي ≈ {words:,} كلمة)", ""]
    out += [f"• {s}" for s in ar]
    figs = key_figures(text, set(sents))
    if figs:
        out += ["", "📌 أرقام ومبالغ بارزة:"] + [f"• {s}" for s in translate_lines(figs)]
    return "\n".join(out)


# ──────────────────────────── استخراج النصوص ────────────────────────────
def extract_text(data, name):
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    with tempfile.NamedTemporaryFile(suffix="." + (ext or "bin"), delete=False) as f:
        f.write(data)
        path = f.name
    try:
        if ext == "pdf":
            from pypdf import PdfReader
            reader = PdfReader(path)
            pages = [(p.extract_text() or "") for p in reader.pages[:300]]
            text = "\n\n".join(pages)
            if len(text.strip()) < 50:
                raise RuntimeError("ملف PDF ممسوح ضوئياً (صور بلا نص). أرسل صفحاته كصور ليتم استخراج النص منها.")
            return text
        if ext == "docx":
            import docx
            d = docx.Document(path)
            parts = [p.text for p in d.paragraphs]
            for tb in d.tables:
                for row in tb.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)
        if ext == "pptx":
            from pptx import Presentation
            prs, parts = Presentation(path), []
            for i, slide in enumerate(prs.slides, 1):
                parts.append(f"\nشريحة {i}")
                for sh in slide.shapes:
                    if sh.has_text_frame:
                        parts.append(sh.text_frame.text)
            return "\n".join(parts)
        if ext in ("xlsx", "xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"\nورقة: {ws.title}")
                for n, row in enumerate(ws.iter_rows(values_only=True)):
                    if n > 2000:
                        break
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        if ext in ("html", "htm"):
            soup = BeautifulSoup(data, "lxml")
            for t in soup(["script", "style", "noscript"]):
                t.decompose()
            return soup.get_text("\n")
        if ext in ("png", "jpg", "jpeg", "webp", "bmp", "tif", "tiff"):
            return ocr_image(path)
        if ext in ("doc", "xls", "ppt"):
            raise RuntimeError("الصيغ القديمة (doc/xls/ppt) غير مدعومة، احفظ الملف بصيغة docx/xlsx/pptx.")
        for enc in ("utf-8", "cp1256", "latin-1"):
            try:
                return data.decode(enc)
            except Exception:
                continue
        raise RuntimeError("صيغة الملف غير مدعومة.")
    finally:
        try:
            os.remove(path)
        except Exception:
            pass


def ocr_image(path):
    try:
        import pytesseract
        from PIL import Image
        img = Image.open(path)
        try:
            txt = pytesseract.image_to_string(img, lang="ara+eng")
        except Exception:
            txt = pytesseract.image_to_string(img)
        if not txt.strip():
            raise RuntimeError("لم أجد نصاً داخل الصورة.")
        return txt
    except ImportError:
        raise RuntimeError("OCR غير مثبّت على السيرفر.")


def fetch_article(url):
    r = requests.get(url, headers=SEC_HEADERS if "sec.gov" in url else WEB_HEADERS, timeout=30)
    r.raise_for_status()
    soup = BeautifulSoup(r.content, "lxml")
    title = (soup.title.string or url).strip() if soup.title else url
    for t in soup(["script", "style", "noscript", "nav", "footer", "aside", "form", "header"]):
        t.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    paras = [p.get_text(" ", strip=True) for p in root.find_all(["p", "h1", "h2", "h3", "li"])]
    text = "\n".join(p for p in paras if len(p) > 25)
    if len(text) < 200:
        text = root.get_text("\n", strip=True)
    return title, text


# ──────────────────────────── تحليل الأسهم ────────────────────────────
def num(x):
    try:
        v = float(x)
        return None if (math.isnan(v) or math.isinf(v)) else v
    except Exception:
        return None


def money(x, cur="USD"):
    v = num(x)
    if v is None:
        return "غير متوفر"
    a, sign = abs(v), "-" if v < 0 else ""
    for div, s in ((1e12, "T"), (1e9, "B"), (1e6, "M"), (1e3, "K")):
        if a >= div:
            return f"{sign}{a / div:.2f}{s} {cur}"
    return f"{sign}{a:,.2f} {cur}"


def pct(a, b):
    a, b = num(a), num(b)
    if a is None or b in (None, 0):
        return "—"
    return f"{(a - b) / abs(b) * 100:+.1f}%"


def get_row(df, names):
    try:
        if df is None or df.empty:
            return None
        for n in names:
            if n in df.index:
                s = df.loc[n].dropna()
                if len(s):
                    return s.sort_index(ascending=False)
    except Exception:
        pass
    return None


def find_symbol(text, strict):
    t = (text or "").strip()
    if "tradingview.com" in t.lower():
        m = re.search(r"symbol=(?:[A-Za-z_]+(?:%3A|:))?([A-Za-z0-9\.\-]+)", t) or \
            re.search(r"/symbols/(?:[A-Za-z_]+-)?([A-Za-z0-9\.]+)", t)
        if m:
            return norm_symbol(m.group(1))
    m = re.search(r"\$([A-Za-z]{1,5}(?:[.\-][A-Za-z]{1,2})?)\b", t)
    if m:
        return norm_symbol(m.group(1))
    tick = r"([A-Za-z]{1,5}(?:[.\-][A-Za-z]{1,2})?|\d{4}(?:\.SR)?)"
    if strict:
        m = re.fullmatch(tick, t)
        if m and (m.group(1).isupper() or m.group(1)[0].isdigit()):
            return norm_symbol(m.group(1))
        return None
    m = re.search(r"(?<![\w])" + tick + r"(?![\w])", t)
    return norm_symbol(m.group(1)) if m else None


def norm_symbol(s):
    s = s.upper()
    if re.fullmatch(r"\d{4}", s):
        return s + ".SR"
    m = re.fullmatch(r"([A-Z]{1,5})\.([A-Z])", s)
    return f"{m.group(1)}-{m.group(2)}" if m else s


_SEC_MAP = {}


def sec_cik(sym):
    if not _SEC_MAP:
        r = requests.get("https://www.sec.gov/files/company_tickers.json", headers=SEC_HEADERS, timeout=30)
        for v in r.json().values():
            _SEC_MAP[v["ticker"].upper()] = str(v["cik_str"]).zfill(10)
    return _SEC_MAP.get(sym.upper())


def sec_recent(cik, days=180):
    j = requests.get(f"https://data.sec.gov/submissions/CIK{cik}.json", headers=SEC_HEADERS, timeout=30).json()
    r = j["filings"]["recent"]
    cutoff = (datetime.now(timezone.utc) - timedelta(days=days)).date().isoformat()
    items_list = r.get("items") or []
    out = []
    for i in range(len(r["form"])):
        if r["filingDate"][i] < cutoff:
            break
        acc = r["accessionNumber"][i].replace("-", "")
        out.append({
            "form": r["form"][i],
            "date": r["filingDate"][i],
            "items": items_list[i] if i < len(items_list) else "",
            "url": f"https://www.sec.gov/Archives/edgar/data/{int(cik)}/{acc}/{r['primaryDocument'][i]}",
        })
    return out


ITEM_AR = {
    "1.01": "اتفاقية / عقد جوهري",
    "1.02": "إنهاء اتفاقية",
    "2.02": "نتائج مالية",
    "2.03": "التزام / دين جديد",
    "3.01": "إشعار إدراج (مخالفة / نقص امتثال)",
    "3.02": "إصدار أسهم غير مسجّل (زيادة رأس المال)",
    "3.03": "تعديل حقوق المساهمين",
    "5.02": "تغيير إداري",
    "5.03": "تعديل النظام الأساسي (قد يتضمن تقسيم أو تغيير أسهم)",
    "5.07": "نتائج تصويت المساهمين",
    "7.01": "إفصاح / عرض",
    "8.01": "حدث آخر",
}
CAP_PREFIX = ("S-1", "S-3", "F-1", "F-3", "424B", "S-11")


def filing_label(f):
    form = f["form"]
    if form.startswith(CAP_PREFIX):
        return f"{form} — تسجيل / طرح أوراق مالية (احتمال تخفيف)"
    if form.startswith("8-K"):
        labels = [ITEM_AR[i] for i in f["items"].split(",") if i in ITEM_AR]
        return f"{form} — " + ("، ".join(labels) if labels else "إفصاح")
    if form in ("DEF 14A", "PRE 14A"):
        return f"{form} — بيان تفويض مساهمين (قد يتضمن زيادة أسهم أو تقسيم عكسي)"
    return form


def filing_lines(fl, limit=5):
    return [f"• {f['date']} — {filing_label(f)}\n  {f['url']}" for f in fl[:limit]]


def filing_brief(url):
    try:
        r = requests.get(url, headers=SEC_HEADERS, timeout=30)
        soup = BeautifulSoup(r.content, "lxml")
        text = soup.get_text("\n")
        text = text[text.find("Item"):] if "Item" in text else text
        sents = split_sentences(text)[:3]
        return translate_ar(" ".join(sents))[:450] if sents else ""
    except Exception:
        return ""


def parse_dt(ts):
    try:
        if isinstance(ts, (int, float)):
            return datetime.fromtimestamp(ts, tz=timezone.utc)
        if isinstance(ts, str):
            return datetime.fromisoformat(ts.replace("Z", "+00:00"))
    except Exception:
        pass
    return None


def get_news(t, sym):
    items = []
    try:
        for n in (t.news or []):
            c = n.get("content") or n
            title = c.get("title")
            if not title:
                continue
            link = ((c.get("canonicalUrl") or {}).get("url") or (c.get("clickThroughUrl") or {}).get("url")
                    or c.get("link") or "")
            pub = (c.get("provider") or {}).get("displayName") or c.get("publisher") or ""
            items.append((parse_dt(c.get("pubDate") or c.get("providerPublishTime")), title, pub, link))
    except Exception as e:  # noqa
        log.warning("yf news: %s", e)
    try:
        r = requests.get(f"https://feeds.finance.yahoo.com/rss/2.0/headline?s={sym}&region=US&lang=en-US",
                         headers=WEB_HEADERS, timeout=20)
        for it in ET.fromstring(r.content).iter("item"):
            title = (it.findtext("title") or "").strip()
            if not title:
                continue
            try:
                dt = parsedate_to_datetime(it.findtext("pubDate"))
            except Exception:
                dt = None
            items.append((dt, title, "Yahoo Finance", (it.findtext("link") or "").strip()))
    except Exception as e:  # noqa
        log.warning("rss: %s", e)
    seen, uniq = set(), []
    for it in items:
        k = it[1].lower()[:80]
        if k not in seen:
            seen.add(k)
            uniq.append(it)
    epoch = datetime(1970, 1, 1, tzinfo=timezone.utc)
    uniq.sort(key=lambda x: x[0] or epoch, reverse=True)
    return uniq


NEWS_KW = re.compile(r"contract|agreement|award|order|partner|deal|project|collaborat|acquir|acquisition|"
                     r"license|supply|FDA|approv|launch|grant|reverse split|offering|private placement", re.I)


def analyze_stock(sym):
    import yfinance as yf
    t = yf.Ticker(sym)
    hist = t.history(period="1y", interval="1d")
    if hist is None or hist.empty:
        return [f"❌ لم أجد بيانات للرمز {sym}. تأكد من الرمز (مثال: AAPL أو 2222 للسوق السعودي)."]
    try:
        info = t.info or {}
    except Exception:
        info = {}
    cur = info.get("currency") or "USD"
    name = info.get("longName") or info.get("shortName") or sym
    c = hist["Close"].dropna()
    price = num(c.iloc[-1])
    now = datetime.now(timezone.utc)

    # بيانات SEC (للأسهم الأمريكية)
    filings, sec_note = [], ""
    try:
        cik = sec_cik(sym)
        if cik:
            filings = sec_recent(cik, 180)
        else:
            sec_note = "ℹ️ لا توجد بيانات SEC لهذا الرمز (سهم غير مدرج أمريكياً أو غير مسجّل)."
    except Exception as e:  # noqa
        sec_note = "ℹ️ تعذّر الوصول إلى SEC حالياً."
        log.warning("sec: %s", e)

    sections = []

    # ── الترويسة
    sections.append(
        f"📊 {name} ({sym})\n"
        f"السعر الحالي: {price:.2f} {cur}\n"
        f"القيمة السوقية: {money(info.get('marketCap'), cur)}\n"
        f"البورصة: {info.get('exchange') or '—'}  ·  الدولة: {info.get('country') or '—'}"
    )

    # ── 1) النشاط
    try:
        cat = " / ".join(x for x in (info.get("sector"), info.get("industry")) if x)
        desc_sents = split_sentences(info.get("longBusinessSummary") or "")[:2]
        tr = translate_lines([cat, " ".join(desc_sents)]) if (cat or desc_sents) else ["", ""]
        lines = []
        if tr[0]:
            lines.append(f"• القطاع: {tr[0]}")
        if tr[1]:
            lines.append(f"• نبذة: {tr[1]}")
        if not lines:
            lines.append("• لا تتوفر بيانات النشاط.")
        sections.append("🏢 النشاط\n" + "\n".join(lines))
    except Exception as e:  # noqa
        log.warning("activity: %s", e)

    # ── 2) نتائج الربع
    try:
        qi = t.quarterly_income_stmt
        rev = get_row(qi, ["Total Revenue", "Operating Revenue"])
        ni = get_row(qi, ["Net Income", "Net Income Common Stockholders"])
        eps = get_row(qi, ["Diluted EPS", "Basic EPS"])
        lines = []
        if rev is not None or ni is not None:
            base = (rev if rev is not None else ni)
            lines.append(f"آخر ربع معلن: {base.index[0].strftime('%Y-%m-%d')}")

            def cmp(label, s, is_money=True):
                v = s.iloc[0]
                txt = f"• {label}: {money(v, cur) if is_money else f'{v:.2f}'}"
                if len(s) > 1:
                    txt += f"  (ربع سابق {pct(v, s.iloc[1])})"
                if len(s) > 4:
                    txt += f"  (نفس الربع العام الماضي {pct(v, s.iloc[4])})"
                return txt
            if rev is not None:
                lines.append(cmp("الإيرادات", rev))
            if ni is not None:
                lines.append(cmp("صافي الربح / الخسارة", ni))
            if eps is not None:
                lines.append(cmp("ربحية السهم", eps, False))
        else:
            lines.append("• لا تتوفر نتائج ربعية.")
        try:
            ed = t.get_earnings_dates(limit=6)
            rep = ed[ed["Reported EPS"].notna()].iloc[0]
            lines.append(f"• مفاجأة النتائج: الفعلي {rep['Reported EPS']} مقابل المتوقع {rep['EPS Estimate']}")
        except Exception:
            pass
        try:
            nxt = (t.calendar or {}).get("Earnings Date")
            if nxt:
                lines.append(f"• موعد النتائج القادم: {nxt[0]}")
        except Exception:
            pass
        sections.append("📈 نتائج الربع\n" + "\n".join(lines))
    except Exception as e:  # noqa
        log.warning("results: %s", e)

    # ── 3) المستهدفات
    try:
        lines = []
        tm = num(info.get("targetMeanPrice"))
        if tm:
            lines.append(f"• متوسط مستهدف المحللين: {tm:.2f}  ({pct(tm, price)} عن السعر الحالي)")
            lines.append(f"• أعلى مستهدف: {info.get('targetHighPrice')}  ·  أدنى: {info.get('targetLowPrice')}"
                         f"  ·  عدد المحللين: {info.get('numberOfAnalystOpinions')}")
            if info.get("recommendationKey"):
                lines.append(f"• التوصية العامة: {info['recommendationKey']}")
        else:
            lines.append("• لا توجد تغطية محللين (شائع في الأسهم الصغيرة).")
        hi52, lo52 = num(hist["High"].max()), num(hist["Low"].min())
        hi3m = num(hist["High"].tail(63).max())
        lines.append(f"• أعلى 3 أشهر: {hi3m:.2f}  ·  أعلى 52 أسبوع: {hi52:.2f}")
        lines.append(f"• أدنى 52 أسبوع: {lo52:.2f}")
        sections.append("🎯 المستهدفات والمقاومات\n" + "\n".join(lines))
    except Exception as e:  # noqa
        log.warning("targets: %s", e)

    # ── 4) المشاريع والعقود
    news = get_news(t, sym)
    try:
        lines = []
        contracts = [f for f in filings if f["form"].startswith("8-K") and "1.01" in f["items"]]
        if contracts:
            lines.append("عقود واتفاقيات جوهرية (SEC — آخر 6 أشهر):")
            for f in contracts[:3]:
                lines.append(f"• {f['date']} — {f['form']}\n  {f['url']}")
            for f in contracts[:2]:
                b = filing_brief(f["url"])
                if b:
                    lines.append(f"  ↳ {b}")
        kw = [n for n in news if NEWS_KW.search(n[1]) and (not n[0] or now - n[0] < timedelta(days=45))][:4]
        if kw:
            if lines:
                lines.append("")
            lines.append("من العناوين الأخيرة (مشاريع / عقود / شراكات):")
            for (dt, title, pub, link), ar in zip(kw, translate_lines([k[1] for k in kw])):
                lines.append(f"• {ar}\n  {link}")
        if not lines:
            lines.append("• لا توجد عقود أو مشاريع معلنة حديثاً في المصادر المتاحة.")
        sections.append("🤝 المشاريع والعقود\n" + "\n".join(lines))
    except Exception as e:  # noqa
        log.warning("contracts: %s", e)

    # ── 5) الاتجاه العام
    try:
        sma = {n: num(c.rolling(n).mean().iloc[-1]) for n in (20, 50, 200)}
        above = [n for n, v in sma.items() if v and price > v]
        valid = [n for n, v in sma.items() if v]
        if valid and len(above) == len(valid):
            trend = "صاعد 🟢  (فوق جميع المتوسطات)"
        elif not above:
            trend = "هابط 🔴  (تحت جميع المتوسطات)"
        else:
            trend = "مختلط / عرضي 🟡"
        d = c.diff()
        up, dn = d.clip(lower=0), -d.clip(upper=0)
        rs = up.ewm(alpha=1 / 14, adjust=False).mean() / dn.ewm(alpha=1 / 14, adjust=False).mean()
        rsi = num((100 - 100 / (1 + rs)).iloc[-1])

        def chg(n):
            return pct(price, c.iloc[-n - 1]) if len(c) > n else "—"
        vol_avg = num(hist["Volume"].tail(30).mean())
        vol_now = num(hist["Volume"].iloc[-1])
        lines = [
            f"• الاتجاه: {trend}",
            "• المتوسطات: " + "  ·  ".join(f"MA{n}: {v:.2f}" for n, v in sma.items() if v),
            f"• التغير: شهر {chg(21)}  ·  3 أشهر {chg(63)}  ·  6 أشهر {chg(126)}",
        ]
        if rsi is not None:
            state = "تشبع شراء" if rsi > 70 else "تشبع بيع" if rsi < 30 else "محايد"
            lines.append(f"• RSI(14): {rsi:.0f}  ({state})")
        if vol_avg and vol_now:
            lines.append(f"• حجم آخر جلسة: {vol_now / vol_avg:.1f}× متوسط 30 يوماً")
        sections.append("🧭 الاتجاه العام\n" + "\n".join(lines))
    except Exception as e:  # noqa
        log.warning("trend: %s", e)

    # ── 6) رأس المال / توزيعات / تقسيم عكسي
    try:
        lines = []
        bs = t.quarterly_balance_sheet
        if bs is None or bs.empty:
            bs = t.balance_sheet
        sh = get_row(bs, ["Ordinary Shares Number", "Share Issued"])
        if sh is not None and len(sh) > 1:
            old = sh.iloc[min(4, len(sh) - 1)]
            ch = pct(sh.iloc[0], old)
            dil = "  ⚠️ تخفيف ملحوظ" if (num(sh.iloc[0]) or 0) > (num(old) or 1) * 1.1 else ""
            lines.append(f"• الأسهم القائمة: {sh.iloc[0]:,.0f}  (تغير {ch} خلال {min(4, len(sh) - 1)} أرباع){dil}")
        elif info.get("sharesOutstanding"):
            lines.append(f"• الأسهم القائمة: {info['sharesOutstanding']:,.0f}")
        cap = [f for f in filings if f["form"].startswith(CAP_PREFIX)
               or (f["form"].startswith("8-K") and "3.02" in f["items"])]
        if cap:
            lines.append("• طروحات / زيادة رأس مال (SEC — آخر 6 أشهر):")
            lines += ["  " + l.replace("\n", "\n  ") for l in filing_lines(cap, 5)]
        else:
            lines.append("• لا توجد إيداعات طرح أو زيادة رأس مال في آخر 6 أشهر." if filings else "• زيادة رأس المال: غير متاح.")
        # التقسيم
        try:
            sp = t.splits
            if sp is not None and len(sp):
                tz = sp.index.tz
                recent = sp[sp.index > (pd_now(tz) - timedelta(days=5 * 365))]
                for dt, ratio in list(recent.items())[-3:]:
                    r = float(ratio)
                    desc = f"تقسيم عكسي 1:{round(1 / r)}" if r < 1 else f"تقسيم أسهم {round(r, 2)}:1"
                    lines.append(f"• {dt.strftime('%Y-%m-%d')} — {desc}")
            else:
                lines.append("• لا يوجد تقسيم (عادي أو عكسي) خلال 5 سنوات.")
        except Exception:
            pass
        rs_f = [f for f in filings if (f["form"].startswith("8-K") and "5.03" in f["items"])
                or f["form"] in ("DEF 14A", "PRE 14A")]
        if rs_f:
            lines.append("• إيداعات قد ترتبط بتقسيم عكسي أو تعديل أسهم (يُنصح بمراجعة المستند):")
            lines += ["  " + l.replace("\n", "\n  ") for l in filing_lines(rs_f, 3)]
        rs_news = [n for n in news if re.search(r"reverse (stock )?split", n[1], re.I)][:2]
        for n in rs_news:
            lines.append(f"• خبر: {translate_ar(n[1])}\n  {n[3]}")
        warn = [f for f in filings if f["form"].startswith("8-K") and "3.01" in f["items"]]
        if warn:
            lines.append("• ⚠️ إشعار إدراج / امتثال:")
            lines += ["  " + l.replace("\n", "\n  ") for l in filing_lines(warn, 2)]
        # التوزيعات
        dv = t.dividends
        if dv is not None and len(dv):
            last = dv.iloc[-1]
            yield_txt = f"  ·  العائد {info['dividendYield']:.2f}%" if num(info.get("dividendYield")) else ""
            lines.append(f"• آخر توزيع: {last:.4f} {cur} بتاريخ {dv.index[-1].strftime('%Y-%m-%d')}{yield_txt}")
        else:
            lines.append("• لا توجد توزيعات أرباح.")
        sections.append("🏦 رأس المال والتوزيعات والتقسيم\n" + "\n".join(lines))
    except Exception as e:  # noqa
        log.warning("capital: %s", e)

    # ── 7) الخسائر المتراكمة والسيولة
    try:
        bs = t.quarterly_balance_sheet
        if bs is None or bs.empty:
            bs = t.balance_sheet
        re_ = get_row(bs, ["Retained Earnings"])
        lines = []
        if re_ is not None:
            v = re_.iloc[0]
            if v < 0:
                lines.append(f"• خسائر متراكمة: {money(abs(v), cur)}  (بتاريخ {re_.index[0].strftime('%Y-%m-%d')})")
            else:
                lines.append(f"• أرباح محتجزة: {money(v, cur)}")
        else:
            lines.append("• الأرباح المحتجزة غير متاحة في المصدر.")
        qi = t.quarterly_income_stmt
        ni = get_row(qi, ["Net Income", "Net Income Common Stockholders"])
        if ni is not None and len(ni) >= 4:
            lines.append(f"• صافي آخر 4 أرباع: {money(ni.iloc[:4].sum(), cur)}")
        cash = get_row(bs, ["Cash Cash Equivalents And Short Term Investments", "Cash And Cash Equivalents"])
        ocf = get_row(t.quarterly_cashflow, ["Operating Cash Flow"])
        if cash is not None:
            lines.append(f"• النقد المتاح: {money(cash.iloc[0], cur)}")
            if ocf is not None and len(ocf) >= 4 and ocf.iloc[:4].sum() < 0:
                monthly = -ocf.iloc[:4].sum() / 12
                months = cash.iloc[0] / monthly if monthly else 0
                warn = "  ⚠️ أقل من 12 شهراً" if months < 12 else ""
                lines.append(f"• حرق نقدي تقريبي: {money(monthly, cur)} / شهر  ←  يكفي نحو {months:.1f} شهر{warn}")
        sections.append("📉 الخسائر المتراكمة والسيولة\n" + "\n".join(lines))
    except Exception as e:  # noqa
        log.warning("deficit: %s", e)

    # ── 8) الأخبار
    try:
        today = [n for n in news if n[0] and now - n[0] < timedelta(hours=24)][:6]
        shown, title = (today, "📰 أخبار اليوم") if today else (news[:6], "📰 آخر الأخبار (لا يوجد خلال 24 ساعة)")
        if shown:
            ar = translate_lines([n[1] for n in shown])
            body = []
            for (dt, _, pub, link), a in zip(shown, ar):
                stamp = dt.strftime("%m-%d %H:%M UTC") if dt else ""
                body.append(f"• {a}\n  {pub}  {stamp}\n  {link}")
            sections.append(title + "\n" + "\n".join(body))
        else:
            sections.append("📰 الأخبار\n• لا توجد أخبار متاحة.")
    except Exception as e:  # noqa
        log.warning("news: %s", e)

    if sec_note:
        sections.append(sec_note)
    sections.append("⚠️ المعلومات للاطلاع فقط، وليست توصية بالبيع أو الشراء.")
    return ["\n\n".join(sections)]


def pd_now(tz):
    import pandas as pd
    return pd.Timestamp.now(tz=tz) if tz else pd.Timestamp.now()


# ──────────────────────────── تحليل ملفات SEC (مستوى متخذ القرار) ────────────────────────────
SEC_DOC_RE = re.compile(r"/data/(\d+)/(\d{18})/")

FORM_AR = {
    "3": "ملكية مبدئية لداخلي (Form 3)", "4": "تعاملات داخلية (Form 4)", "5": "تعاملات داخلية سنوية (Form 5)",
    "SCHEDULE 13G": "ملكية ≥5% سلبية (13G)", "SCHEDULE 13D": "ملكية ≥5% بنية نشطة (13D)",
    "SC 13G": "ملكية ≥5% سلبية (13G)", "SC 13D": "ملكية ≥5% بنية نشطة (13D)",
    "8-K": "حدث جوهري (8-K)", "6-K": "إفصاح شركة أجنبية (6-K)", "10-Q": "تقرير ربع سنوي (10-Q)",
    "10-K": "تقرير سنوي (10-K)", "20-F": "تقرير سنوي أجنبي (20-F)", "DEF 14A": "بيان تفويض مساهمين",
}
ITEM_TAKEAWAY = {
    "1.01": "اتفاقية جوهرية — راجع قيمة العقد وطرفه ومدته مقارنة بإيرادات الشركة.",
    "1.02": "إنهاء اتفاقية — سلبي إذا كانت مصدر دخل رئيسي.",
    "1.03": "إفلاس / إعسار — خطر شديد.",
    "2.02": "نتائج مالية — قارن الإيرادات والربحية بالتوقعات والربع السابق.",
    "2.03": "التزام أو دين جديد — راجع الفائدة والضمانات وقابلية التحويل لأسهم.",
    "2.04": "تسريع التزام أو تعثر — خطر سيولة.",
    "2.05": "إعادة هيكلة / تسريح — قد يخفض التكاليف لكنه يعكس ضغطاً تشغيلياً.",
    "2.06": "انخفاض قيمة أصول — أثر محاسبي سلبي.",
    "3.01": "إشعار إدراج — خطر شطب، وغالباً ما يتبعه تقسيم عكسي.",
    "3.02": "إصدار أسهم غير مسجّل — تخفيف مباشر؛ قارن سعر الإصدار بسعر السوق.",
    "3.03": "تعديل حقوق حملة الأسهم.",
    "4.01": "تغيير المحاسب القانوني — راقب الأسباب.",
    "4.02": "عدم الاعتماد على قوائم سابقة (إعادة بيان) — إشارة خطر عالية.",
    "5.01": "تغيير في السيطرة.",
    "5.02": "تغيير في الإدارة أو المجلس — راجع المنصب وسبب المغادرة.",
    "5.03": "تعديل النظام الأساسي — قد يتضمن تقسيماً عكسياً أو زيادة أسهم مصرّح بها.",
    "5.07": "نتائج تصويت المساهمين — قد تتضمن الموافقة على زيادة أسهم أو تقسيم.",
    "7.01": "إفصاح (عرض / بيان).",
    "8.01": "حدث آخر — اقرأ التفاصيل.",
}
TX_AR = {"P": "شراء من السوق المفتوح", "S": "بيع في السوق المفتوح", "A": "منحة/مكافأة أسهم",
         "M": "ممارسة خيار/مشتق", "F": "سداد ضريبة/سعر الممارسة بأسهم", "G": "هبة", "D": "إعادة للشركة",
         "C": "تحويل مشتق", "X": "ممارسة مشتق", "J": "أخرى", "W": "وصية/ميراث", "I": "عملية تقديرية"}
HOLDER_TYPE = {"IA": "مدير استثمار", "BD": "وسيط/تاجر", "IN": "فرد", "CO": "شركة", "HC": "شركة قابضة",
               "PN": "شراكة", "OO": "أخرى", "IC": "شركة تأمين", "EP": "صندوق تقاعد", "FI": "مؤسسة مالية أجنبية"}

RISK_PATTERNS = [
    ("استمرارية المنشأة (Going Concern)", r"substantial doubt|going concern"),
    ("تقسيم عكسي", r"reverse (?:stock |share )?split"),
    ("خطر الشطب/عدم امتثال إدراج", r"minimum bid|delist|listing (?:rule|deficiency)|deficiency notice"),
    ("تخفيف: طرح أسهم/وارنتات", r"private placement|registered direct|at-the-market|securities purchase agreement|pre-funded warrant|public offering"),
    ("دين تحويلي", r"convertible (?:promissory )?(?:note|debenture)"),
    ("تعثر/إفلاس", r"event of default|chapter 11|chapter 7|bankruptcy|insolven"),
    ("إعادة بيان/ضعف رقابة داخلية", r"non-reliance|restate|material weakness"),
    ("استقالة/إنهاء مسؤول", r"resign|stepped down|terminated .{0,30}(?:employment|agreement)"),
]
POS_PATTERNS = [
    ("عقد/طلبية/ترسية", r"awarded|purchase order|(?:entered into|signed) (?:a |an )?(?:\w+ ){0,3}(?:contract|agreement)"),
    ("موافقة/ترخيص تنظيمي", r"FDA (?:approval|clearance)|received (?:regulatory )?approval|clearance from"),
    ("نمو قياسي", r"record (?:revenue|quarter|sales|results)"),
    ("استحواذ/شراكة استراتيجية", r"acquisition of|strategic (?:partnership|investment)|definitive agreement"),
    ("إعادة شراء/توزيعات", r"share repurchase|buyback|special dividend"),
]
OFFER_FIELDS = [
    ("حجم/شروط الطرح", r"[^.]{0,80}\b(?:offering|offer|selling|sell)\b[^.]{0,40}\bup to\b[^.]{0,200}\."),
    ("السعر (طرح/تنفيذ/تحويل)", r"[^.]{0,60}(?:public offering price|offering price|exercise price|conversion price|purchase price)[^.]{0,180}\."),
    ("الوارنتات", r"[^.]{0,60}\bwarrants?\b[^.]{0,200}\."),
    ("العوائد", r"[^.]{0,60}(?:gross|net) proceeds[^.]{0,200}\."),
    ("الأسهم بعد الطرح", r"[^.]{0,60}outstanding\s+(?:immediately\s+)?after[^.]{0,200}\."),
    ("وكيل الطرح/المكتتب", r"[^.]{0,40}(?:placement agent|underwriter|sales agent)[^.]{0,160}\."),
    ("استخدام العوائد", r"use of proceeds[:\.\s]+[^.]{20,300}\."),
]
XB_METRICS = [
    ("الإيرادات", ["Revenues", "RevenueFromContractWithCustomerExcludingAssessedTax", "SalesRevenueNet"], "flow"),
    ("صافي الربح/الخسارة", ["NetIncomeLoss", "ProfitLoss"], "flow"),
    ("الربح التشغيلي", ["OperatingIncomeLoss"], "flow"),
    ("ربحية السهم", ["EarningsPerShareDiluted", "EarningsPerShareBasicAndDiluted", "EarningsPerShareBasic"], "flow"),
    ("التدفق النقدي التشغيلي", ["NetCashProvidedByUsedInOperatingActivities"], "flow"),
    ("النقد", ["CashAndCashEquivalentsAtCarryingValue", "CashCashEquivalentsRestrictedCashAndRestrictedCashEquivalents"], "inst"),
    ("إجمالي الالتزامات", ["Liabilities"], "inst"),
    ("حقوق المساهمين", ["StockholdersEquity", "StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest"], "inst"),
    ("الخسائر المتراكمة (سالب) / الأرباح المحتجزة", ["RetainedEarningsAccumulatedDeficit"], "inst"),
]


def sec_get(url, timeout=60):
    time.sleep(0.2)
    r = requests.get(url, headers=SEC_HEADERS, timeout=timeout)
    r.raise_for_status()
    return r.content


def _ln(tag):
    return tag.rsplit("}", 1)[-1]


def xml_find(root, *subs):
    out = []
    for e in root.iter():
        n = _ln(e.tag).lower()
        if all(s in n for s in subs) and (e.text or "").strip():
            out.append(e.text.strip())
    return out


def xml_one(root, *subs):
    v = xml_find(root, *subs)
    return v[0] if v else ""


def xv(el, name):
    for e in el.iter():
        if _ln(e.tag) == name:
            for c in e.iter():
                if _ln(c.tag) == "value" and (c.text or "").strip():
                    return c.text.strip()
            if (e.text or "").strip():
                return e.text.strip()
    return ""


def html_text(raw):
    soup = BeautifulSoup(raw, "lxml")
    for t in soup(["script", "style", "noscript", "ix:header"]):
        t.decompose()
    return soup.get_text("\n")


def flat(text):
    return re.sub(r"\s+", " ", text or "")


def scan_flags(text, patterns, width=110):
    f = flat(text)
    out = []
    for label, pat in patterns:
        m = re.search(pat, f, re.I)
        if m:
            s, e = max(0, m.start() - width), min(len(f), m.end() + width)
            out.append((label, f[s:e].strip(), len(re.findall(pat, f, re.I))))
    return out


def flags_lines(text):
    lines = []
    risk, pos = scan_flags(text, RISK_PATTERNS), scan_flags(text, POS_PATTERNS)
    if risk:
        tr = translate_lines([x[1] for x in risk])
        lines.append("\n🚩 إشارات خطر:")
        lines += [f"• {r[0]} (تكرر {r[2]}×):\n  «{t}»" for r, t in zip(risk, tr)]
    if pos:
        tr = translate_lines([x[1] for x in pos])
        lines.append("\n✅ إشارات إيجابية:")
        lines += [f"• {p[0]} (تكرر {p[2]}×):\n  «{t}»" for p, t in zip(pos, tr)]
    return lines


def sec_meta(cik, acc):
    dashed = f"{acc[:10]}-{acc[10:12]}-{acc[12:]}"
    meta = {}
    try:
        j = requests.get(f"https://data.sec.gov/submissions/CIK{int(cik):010d}.json",
                         headers=SEC_HEADERS, timeout=30).json()
        meta["name"], meta["tickers"] = j.get("name", ""), j.get("tickers", [])
        r = j["filings"]["recent"]
        if dashed in r["accessionNumber"]:
            i = r["accessionNumber"].index(dashed)
            meta["form"], meta["date"] = r["form"][i], r["filingDate"][i]
            meta["items"] = (r.get("items") or [""] * (i + 1))[i]
    except Exception as e:  # noqa
        log.warning("sec_meta: %s", e)
    return meta


def last_price(sym):
    try:
        import yfinance as yf
        fi = yf.Ticker(sym).fast_info
        return num(fi["last_price"]), num(fi["market_cap"])
    except Exception:
        return None, None


def ticker_by_cik(cik):
    try:
        sec_cik("AAPL")  # تحميل الخريطة
        z = str(int(cik)).zfill(10)
        for k, v in _SEC_MAP.items():
            if v == z:
                return k
    except Exception:
        pass
    return None


# ── 13D / 13G
def analyze_13dg(root, raw_txt, meta):
    form = (meta.get("form") or xml_one(root, "submissiontype") or "13G").upper()
    is_d = "13D" in form
    issuer, cusip = xml_one(root, "issuername"), xml_one(root, "cusip")
    icik, cls, event = xml_one(root, "issuercik"), xml_one(root, "securitiesclasstitle"), xml_one(root, "eventdate")
    names = xml_find(root, "reportingpersonname")
    pcts = [num(x.replace("%", "")) for x in xml_find(root, "classpercent")]
    shs = [num(x.replace(",", "")) for x in xml_find(root, "aggregate", "shares")] or \
          [num(x.replace(",", "")) for x in xml_find(root, "aggregate", "amount")]
    types = xml_find(root, "typeofreportingperson")
    low = raw_txt.lower()
    basis = ("مؤسسة مالية مؤهلة (13d-1(b))" if "13d-1(b)" in low else
             "مستثمر سلبي (13d-1(c))" if "13d-1(c)" in low else
             "مالك قبل الطرح/معفى (13d-1(d))" if "13d-1(d)" in low else "")
    ticker = ticker_by_cik(icik) if icik else None
    price, mcap = last_price(ticker) if ticker else (None, None)

    L = [
        f"🏢 الشركة المُصدِرة: {issuer or '—'}" + (f" ({ticker})" if ticker else "") +
        (f"  ·  CUSIP {cusip}" if cusip else ""),
        f"📦 فئة الأوراق: {cls or '—'}  ·  تاريخ الحدث: {event or '—'}" +
        (f"\n🏷 أساس الإفصاح: {basis}" if basis else ""),
        "\n👤 المُفصحون:",
    ]
    top = None
    for i, n in enumerate(names or ["—"]):
        p = pcts[i] if i < len(pcts) else None
        s = shs[i] if i < len(shs) else None
        t = HOLDER_TYPE.get(types[i], types[i]) if i < len(types) else ""
        row = f"• {n}" + (f"  [{t}]" if t else "")
        if s is not None:
            row += f": {s:,.0f} سهم"
        if p is not None:
            row += f"  =  {p:.2f}% من الفئة"
            top = max(top or 0, p)
        if s is not None and price:
            row += f"  ≈  {money(s * price)}"
        L.append(row)
    if mcap:
        L.append(f"\nالقيمة السوقية الحالية للشركة ≈ {money(mcap)}")
    L.append("\n🎯 الدلالة:")
    if is_d:
        L.append("• 13D: ملكية ≥5% مع احتمال نية التأثير أو السيطرة (نشاط مساهم) — أولوية عالية.")
    else:
        L.append("• 13G: استثمار سلبي بلا نية سيطرة. يدل على اهتمام مؤسسي دون تدخل في الإدارة.")
    if top is not None:
        if "/A" in form and top < 5:
            L.append("• ⚠️ النسبة أقل من 5% في تعديل — غالباً خروج أو تخفيض تحت حد الإفصاح (ضغط بيعي محتمل).")
        elif top >= 20:
            L.append("• مالك بنسبة ≥20%: تأثير جوهري على القرارات والسيولة.")
        elif top >= 10:
            L.append("• مالك بنسبة ≥10%: مساهم كبير (يخضع أيضاً لإفصاح Form 4).")
        else:
            L.append("• النسبة قريبة من حد الإفصاح 5%.")
    if "/A" in form:
        L.append("• هذا تعديل: قارن بالإيداع السابق لنفس الجهة لمعرفة إن زادت الحصة أو نقصت.")
    if is_d:
        for key, lab in (("item3", "مصدر التمويل"), ("item4", "الغرض من الصفقة / النوايا")):
            for e in root.iter():
                if _ln(e.tag).lower() == key:
                    t = " ".join(x.strip() for x in e.itertext() if x.strip())
                    if len(t) > 40:
                        L.append(f"\n📝 {lab}:\n{translate_ar(t[:1500])}")
                        break
    return L


# ── Form 3/4/5
def analyze_ownership(root, raw_txt, meta):
    issuer, sym = xml_one(root, "issuername"), xml_one(root, "issuertradingsymbol")
    owners = []
    for ro in root.iter():
        if _ln(ro.tag) == "reportingOwner":
            rel = []
            if xv(ro, "isDirector") in ("1", "true"):
                rel.append("عضو مجلس إدارة")
            if xv(ro, "isOfficer") in ("1", "true"):
                rel.append(xv(ro, "officerTitle") or "مسؤول تنفيذي")
            if xv(ro, "isTenPercentOwner") in ("1", "true"):
                rel.append("مالك +10%")
            owners.append(f"{xv(ro, 'rptOwnerName')} ({'، '.join(rel) or 'مرتبط'})")
    L = [
        f"🏢 الشركة: {issuer or '—'}" + (f" ({sym})" if sym else ""),
        "👤 الداخلي: " + "  ·  ".join(owners),
    ]
    rows = []
    for tx in root.iter():
        if _ln(tx.tag) == "nonDerivativeTransaction":
            rows.append(dict(title=xv(tx, "securityTitle"), date=xv(tx, "transactionDate"),
                             code=xv(tx, "transactionCode"), sh=num(xv(tx, "transactionShares")),
                             px=num(xv(tx, "transactionPricePerShare")),
                             after=num(xv(tx, "sharesOwnedFollowingTransaction")),
                             own=xv(tx, "directOrIndirectOwnership")))
    if rows:
        L.append("\n💱 المعاملات:")
        for r in rows[:10]:
            val = f"  ≈  {money((r['sh'] or 0) * r['px'])}" if r["sh"] and r["px"] else ""
            after_txt = f"  ·  الرصيد بعدها {r['after']:,.0f}" if r["after"] is not None else ""
            indir = "  (غير مباشر)" if r["own"] == "I" else ""
            L.append(
                f"• {r['date']} — {TX_AR.get(r['code'], r['code'])}: "
                f"{r['sh'] or 0:,.0f} سهم"
                + (f" @ {r['px']:.2f}" if r["px"] else "")
                + val + after_txt + indir
            )
    holds = [xv(h, "sharesOwnedFollowingTransaction") for h in root.iter() if _ln(h.tag) == "nonDerivativeHolding"]
    if holds and not rows:
        L.append("\n📦 الحيازات المفصح عنها: " + "، ".join(f"{num(h) or 0:,.0f}" for h in holds[:5]) + " سهم")
    der = [t for t in root.iter() if _ln(t.tag) == "derivativeTransaction"]
    for d in der[:4]:
        L.append(
            f"• مشتق / خيار: {xv(d, 'securityTitle')} — "
            f"{TX_AR.get(xv(d, 'transactionCode'), xv(d, 'transactionCode'))}"
            f"  ·  سعر التنفيذ {xv(d, 'conversionOrExercisePrice') or '—'}"
        )
    buys = sum((r["sh"] or 0) * (r["px"] or 0) for r in rows if r["code"] == "P")
    sells = sum((r["sh"] or 0) * (r["px"] or 0) for r in rows if r["code"] == "S")
    sold_sh = sum(r["sh"] or 0 for r in rows if r["code"] == "S")
    after = next((r["after"] for r in reversed(rows) if r["code"] == "S" and r["after"] is not None), None)
    plan = "10b5-1" in raw_txt.lower()
    L.append("\n🎯 الدلالة:")
    if buys:
        L.append(
            f"• ✅ شراء من السوق المفتوح بمبلغ ≈ {money(buys)} بواسطة داخلي — "
            "من أقوى إشارات الثقة (قيّم الحجم نسبةً لحصته)."
        )
    if sells:
        pct_sold = (
            f" (≈{sold_sh / (after + sold_sh) * 100:.1f}% من حصته)"
            if after is not None and (after + sold_sh) else ""
        )
        plan_txt = (
            "نُفّذ ضمن خطة 10b5-1 مسبقة → دلالة أضعف."
            if plan else
            "لا يظهر أنه ضمن خطة مسبقة → دلالة أقوى."
        )
        L.append(f"• 🔻 بيع بمبلغ ≈ {money(sells)}{pct_sold}. {plan_txt}")
    if not buys and not sells:
        L.append("• حركة روتينية (منح / ممارسة / ضريبة / حيازة) وليست قراراً استثمارياً في السوق.")
    return L


# ── نشرات الطرح
def analyze_offering(text, meta):
    f = flat(text)
    L = ["💼 بنود الطرح المستخرجة (يُنصح بمراجعة المستند للتفاصيل الدقيقة):"]
    found = []
    for label, pat in OFFER_FIELDS:
        m = re.search(pat, f, re.I)
        if m:
            found.append((label, m.group(0).strip()[:380]))
    if found:
        tr = translate_lines([x[1] for x in found])
        L += [f"• {a[0]}:\n  {t}" for a, t in zip(found, tr)]
    else:
        L.append("• لم أتمكن من استخراج بنود الطرح آلياً.")
    L.append(
        "\n🎯 الدلالة:\n"
        "• كل إصدار جديد يخفّف حصص المساهمين الحاليين.\n"
        "• قارن سعر الطرح بسعر السوق الحالي.\n"
        "• قارن حجم الطرح بالقيمة السوقية.\n"
        "• راقب وجود وارنتات (تخفيف إضافي محتمل مستقبلاً)."
    )
    return L + flags_lines(text)


# ── 10-Q / 10-K عبر XBRL
def _d(s):
    return datetime.strptime(s, "%Y-%m-%d")


def xbrl_lines(cik, acc):
    dashed = f"{acc[:10]}-{acc[10:12]}-{acc[12:]}"
    j = requests.get(f"https://data.sec.gov/api/xbrl/companyfacts/CIK{int(cik):010d}.json",
                     headers=SEC_HEADERS, timeout=90).json()
    us = j.get("facts", {}).get("us-gaap", {})
    L, got = [], {}
    for label, tags, kind in XB_METRICS:
        for tag in tags:
            node = us.get(tag)
            if not node:
                continue
            unit = "USD" if "USD" in node["units"] else ("USD/shares" if "USD/shares" in node["units"] else None)
            rows = [r for r in node["units"].get(unit, []) if r.get("accn") == dashed] if unit else []
            if not rows:
                continue
            if kind == "inst":
                ends = sorted({r["end"] for r in rows}, reverse=True)
                v = next(r["val"] for r in rows if r["end"] == ends[0])
                pv = next((r["val"] for r in rows if len(ends) > 1 and r["end"] == ends[1]), None)
                got[label] = (v, pv, None, ends[0], "عن نهاية الفترة السابقة")
            else:
                flows = [r for r in rows if r.get("start")]
                if not flows:
                    continue
                ce = max(r["end"] for r in flows)
                c = min((r for r in flows if r["end"] == ce), key=lambda r: (_d(r["end"]) - _d(r["start"])).days)
                dur = (_d(c["end"]) - _d(c["start"])).days
                pv = None
                for r in flows:
                    gap = (_d(c["end"]) - _d(r["end"])).days
                    if 330 <= gap <= 400 and abs((_d(r["end"]) - _d(r["start"])).days - dur) <= 20:
                        pv = r["val"]
                        break
                got[label] = (c["val"], pv, dur, ce, "عن نفس الفترة من العام الماضي")
            break
    if not got:
        return ["• لا تتوفر بيانات XBRL لهذا الإيداع."]
    any_end = next(iter(got.values()))[3]
    dur = next((g[2] for g in got.values() if g[2]), None)
    L.append(f"📅 الفترة المنتهية: {any_end}" + (f"  (مدة ≈ {round(dur / 30.4)} شهر)" if dur else ""))
    for label, (v, pv, d, end, cmp_txt) in got.items():
        is_eps = "ربحية" in label
        val = f"{v:.2f}$" if is_eps else money(v)
        row = f"• {label}: {val}"
        if pv is not None:
            row += f"  ({cmp_txt}: {pct(v, pv)})"
        L.append(row)
    try:
        dei = j["facts"]["dei"]["EntityCommonStockSharesOutstanding"]["units"]["shares"]
        sh = [r for r in dei if r.get("accn") == dashed]
        if sh:
            L.append(f"• الأسهم القائمة (غلاف التقرير): {sh[-1]['val']:,.0f}")
    except Exception:
        pass
    cash, ocf = got.get("النقد"), got.get("التدفق النقدي التشغيلي")
    L.append("\n🎯 الدلالة:")
    if ocf and ocf[0] < 0 and cash and ocf[2]:
        burn = -ocf[0] / (ocf[2] / 30.4)
        months = cash[0] / burn if burn else 0
        warn = "  ⚠️ أقل من 12 شهراً — احتمال طرح أسهم أو دين قريب" if months < 12 else ""
        L.append(f"• حرق نقدي ≈ {money(burn)} / شهر  ←  النقد يكفي نحو {months:.1f} شهر{warn}")
    eq = got.get("حقوق المساهمين")
    if eq and eq[0] < 0:
        L.append("• ⚠️ حقوق المساهمين سالبة: الالتزامات تفوق الأصول (خطر إدراج / استمرارية).")
    ni = got.get("صافي الربح/الخسارة")
    if ni and ni[0] < 0:
        deeper = " (والخسارة أعمق من العام الماضي)" if ni[1] is not None and ni[0] < ni[1] else ""
        L.append(f"• الشركة خاسرة في هذه الفترة{deeper}.")
    elif ni:
        L.append("• الشركة رابحة في هذه الفترة.")
    return L


# ── نقطة الدخول
def analyze_sec_filing(url):
    m = SEC_DOC_RE.search(url)
    if not m:
        return ["⚠️ لم أتعرف على رابط ملف SEC (يجب أن يكون من /Archives/edgar/data/...)."]
    cik, acc = m.group(1), m.group(2)
    meta = sec_meta(cik, acc)
    raw = sec_get(re.sub(r"/xsl[^/]+/", "/", url))
    txt = raw.decode("utf-8", "ignore")
    head = txt[:4000].lower()
    form = (meta.get("form") or "").upper()
    root = None
    if "<ownershipdocument" in head or "<edgarsubmission" in head:
        try:
            root = ET.fromstring(raw)
        except ET.ParseError:
            root = None
        if not form:
            form = "4" if "<ownershipdocument" in head else "SCHEDULE 13G"
    if not form:
        mm = re.search(r"FORM\s+(8-K|10-Q|10-K|6-K|20-F|S-1|S-3|F-1|F-3|424B\d)", html_text(raw)[:6000], re.I)
        form = mm.group(1).upper() if mm else "OTHER"

    base = form.split("/")[0]
    label = next((v for k, v in FORM_AR.items() if base == k or base.startswith(k + " ")), form)
    name = meta.get("name") or ""
    tick = ", ".join(meta.get("tickers") or [])
    L = [f"📄 {name}" + (f" ({tick})" if tick else ""),
         f"النوع: {label} | الإيداع: {meta.get('date') or '—'}", ""]

    try:
        if root is not None and base in ("3", "4", "5"):
            L += analyze_ownership(root, txt, meta)
        elif root is not None and ("13D" in form or "13G" in form):
            L += analyze_13dg(root, txt, meta)
        else:
            text = html_text(raw)
            if base in ("424B1", "424B2", "424B3", "424B4", "424B5", "424B7", "424B8", "S-1", "S-3", "F-1", "F-3", "S-11"):
                L += analyze_offering(text, meta)
            elif base in ("10-Q", "10-K"):
                L += xbrl_lines(cik, acc)
                L += flags_lines(text)
            elif base in ("8-K", "6-K"):
                L += analyze_8k(text, cik, acc, meta)
            else:
                sents = summarize_extractive(text, 8)
                L.append("📝 أهم ما ورد:")
                L += [f"• {s}" for s in translate_lines(sents)]
                L += flags_lines(text)
    except Exception as e:  # noqa
        log.exception("sec analyze")
        L.append(f"⚠️ تعذّر التحليل المتخصص ({str(e)[:120]}).")
    L += ["", f"🔗 {url}", "⚠️ استنتاجات آلية من نص الإيداع — للاطلاع فقط وليست توصية."]
    return ["\n".join(L)]


def analyze_8k(text, cik, acc, meta):
    pat = re.compile(r"(?im)^\s*item\s+(\d\.\d{2})\b[\.\s:\-–—]*(.*)$")
    ms = list(pat.finditer(text))
    items = {}
    for i, mm in enumerate(ms):
        end = ms[i + 1].start() if i + 1 < len(ms) else len(text)
        sect = re.split(r"(?i)\n\s*signatures?\s*\n", text[mm.end():end])[0]
        body = (mm.group(2) + "\n" + sect).strip()
        if mm.group(1) not in items or len(body) > len(items[mm.group(1)]):
            items[mm.group(1)] = body
    codes = [c for c in items if c != "9.01"] or [c for c in (meta.get("items") or "").split(",") if c and c != "9.01"]
    L = ["🔎 ما حدث (بنود الإفصاح):"]
    L += [f"• {c} — {ITEM_AR.get(c, 'بند')}" for c in codes] or ["• لم أستطع تحديد البنود."]
    pool = ""
    for c in codes:
        body = items.get(c, "")
        if not body:
            continue
        pool += "\n" + body
        sents = summarize_extractive(body, 4) if len(split_sentences(body)) > 4 else split_sentences(body)[:4]
        if sents:
            L.append(f"\n📌 تفاصيل البند {c}:")
            L += [f"• {s}" for s in translate_lines(sents)]
    try:  # البيان الصحفي EX-99
        idx = requests.get(f"https://www.sec.gov/Archives/edgar/data/{int(cik)}/{acc}/index.json",
                           headers=SEC_HEADERS, timeout=30).json()
        names = [x["name"] for x in idx["directory"]["item"]]
        ex = next((n for n in names if re.search(r"(?i)ex-?99.*\.htm", n)), None)
        if ex:
            pr = html_text(sec_get(f"https://www.sec.gov/Archives/edgar/data/{int(cik)}/{acc}/{ex}"))
            pool += "\n" + pr
            sents = summarize_extractive(pr, 6)
            if sents:
                L.append("\n📎 البيان الصحفي المرفق:")
                L += [f"• {s}" for s in translate_lines(sents)]
            figs = key_figures(pr, set(sents), 5)
            if figs:
                L.append("\n💵 أرقام بارزة:")
                L += [f"• {s}" for s in translate_lines(figs)]
    except Exception as e:  # noqa
        log.warning("ex99: %s", e)
    figs = key_figures(pool, set(), 4) if pool and "💵" not in "\n".join(L) else []
    if figs:
        L.append("\n💵 أرقام بارزة:")
        L += [f"• {s}" for s in translate_lines(figs)]
    L += flags_lines(pool)
    take = [f"• {c}: {ITEM_TAKEAWAY[c]}" for c in codes if c in ITEM_TAKEAWAY]
    if take:
        L.append("\n🎯 الدلالة:")
        L += take
    return L


# ──────────────────────────── توجيه الرسائل ────────────────────────────
CMD_ALIAS = {"tr": "tr", "translate": "tr", "sum": "sum", "summary": "sum", "summarize": "sum",
             "stock": "stock", "s": "stock", "analyze": "stock", "start": "help", "help": "help"}
AR_RE = re.compile(r"^(تلخيص|لخص|ترجمة|ترجم|تحليل|حلل)(?:\s+|$)(.*)", re.S)
AR_MAP = {"تلخيص": "sum", "لخص": "sum", "ترجمة": "tr", "ترجم": "tr", "تحليل": "stock", "حلل": "stock"}
URL_RE = re.compile(r"https?://\S+")


def parse_intent(text):
    t = (text or "").strip()
    if t.startswith("/"):
        parts = t.split(None, 1)
        cmd = CMD_ALIAS.get(parts[0][1:].split("@")[0].lower())
        return cmd, (parts[1] if len(parts) > 1 else "")
    m = AR_RE.match(t)
    if m:
        return AR_MAP[m.group(1)], m.group(2).strip()
    return None, t


def media_of(m):
    if not m:
        return None
    if m.get("document"):
        d = m["document"]
        return d["file_id"], d.get("file_name") or "file.bin"
    if m.get("photo"):
        return m["photo"][-1]["file_id"], "photo.jpg"
    return None


def deliver_translation(chat_id, text, title):
    if len(text) > 120000:
        text = text[:120000]
        send(chat_id, "ℹ️ النص طويل جداً، تمت ترجمة أول 120 ألف حرف فقط.")
    out = translate_ar(text)
    if len(out) <= 3500:
        send(chat_id, out)
    else:
        send_doc(chat_id, out, "translation_ar.txt", caption=f"🌐 الترجمة الكاملة: {title}")


def handle(msg):
    chat_id = msg["chat"]["id"]
    if str(chat_id) != CHAT_ID:
        return
    mid = msg.get("message_id")
    text = (msg.get("text") or msg.get("caption") or "").strip()
    cmd, arg = parse_intent(text)
    reply = msg.get("reply_to_message") or {}
    reply_text = (reply.get("text") or reply.get("caption") or "").strip()

    if cmd == "help" or (not text and not media_of(msg) and not media_of(reply)):
        send(chat_id, HELP)
        return

    media = media_of(msg) or (media_of(reply) if cmd in ("tr", "sum") and not arg else None)

    # ─ تحليل سهم (حتى لو وُجدت صورة مرفقة مع «تحليل»)
    sym = None
    if cmd == "stock":
        sym = find_symbol(arg or reply_text, strict=False)
        # إن لم يُذكر الرمز، جرّب استخراجه من نص الصورة (OCR)
        if not sym and media:
            try:
                file_id, fname = media
                send(chat_id, "⏳ جاري البحث عن رمز السهم في الصورة ...", reply_to=mid)
                ocr_txt = extract_text(download_file(file_id), fname)
                sym = find_symbol(ocr_txt, strict=False)
                if not sym:
                    # أنماط شائعة في لقطات الشارت: AB، $AB، NASDAQ:AB
                    m = re.search(
                        r"(?:NASDAQ|NYSE|AMEX|TSE|TSX)[:\s]+([A-Z]{1,5})\b"
                        r"|(?:ticker|symbol|رمز)[:\s]+([A-Z]{1,5})\b"
                        r"|\b([A-Z]{1,5})\s*(?:USD|USD/|@)",
                        ocr_txt, re.I,
                    )
                    if m:
                        sym = norm_symbol(next(g for g in m.groups() if g))
            except Exception as e:  # noqa
                log.warning("ocr symbol: %s", e)
        if not sym:
            send(chat_id, "اكتب رمز السهم مع الأمر، مثال:\nتحليل AAPL\nأو /stock 2222")
            return
        send(chat_id, f"⏳ جاري تحليل {sym} ... (قد يستغرق دقيقة)", reply_to=mid)
        for part in analyze_stock(sym):
            send(chat_id, part)
        return

    if cmd is None and not media:
        sym = find_symbol(text, strict=True)
        if sym:
            send(chat_id, f"⏳ جاري تحليل {sym} ... (قد يستغرق دقيقة)", reply_to=mid)
            for part in analyze_stock(sym):
                send(chat_id, part)
            return

    typing(chat_id)

    # ─ جمع المحتوى (ترجمة / تلخيص)
    title, content = "النص", ""
    if media:
        file_id, fname = media
        send(chat_id, f"⏳ جاري قراءة «{fname}» ...", reply_to=mid)
        content = extract_text(download_file(file_id), fname)
        title = fname
        action = cmd or ("tr" if fname == "photo.jpg" else "sum")
    else:
        src = arg or reply_text
        m = URL_RE.search(src)
        if m and len(src.strip()) < len(m.group(0)) + 20:
            send(chat_id, "⏳ جاري قراءة الرابط ...", reply_to=mid)
            if "sec.gov/Archives" in m.group(0) and cmd != "tr":
                for part in analyze_sec_filing(m.group(0)):
                    send(chat_id, part)
                return
            title, content = fetch_article(m.group(0))
            action = cmd or "sum"
        else:
            content = src
            action = cmd or "tr"

    if len(content.strip()) < 2:
        send(chat_id, "⚠️ لم أجد نصاً لمعالجته.")
        return
    if action == "sum" and len(content) < 400:
        action = "tr"  # النص قصير: الأفضل ترجمته
    if action == "sum":
        content = content[:600000]
        send(chat_id, build_summary(content, title))
    else:
        deliver_translation(chat_id, content, title)


# ──────────────────────────── الحلقة الرئيسية ────────────────────────────
def main():
    log.info("Bot started. Running for %s seconds", RUN_SECONDS)
    deadline = time.time() + RUN_SECONDS
    offset = None
    while time.time() < deadline:
        try:
            r = requests.post(f"{API}/getUpdates", timeout=40, data={
                k: v for k, v in {"offset": offset, "timeout": 25,
                                  "allowed_updates": '["message"]'}.items() if v is not None})
            updates = r.json().get("result", [])
        except Exception as e:  # noqa
            log.warning("getUpdates: %s", e)
            time.sleep(5)
            continue
        for u in updates:
            offset = u["update_id"] + 1
            msg = u.get("message")
            if not msg:
                continue
            try:
                handle(msg)
            except Exception as e:  # noqa
                log.exception("handle error")
                if str(msg["chat"]["id"]) == CHAT_ID:
                    send(msg["chat"]["id"], f"⚠️ حدث خطأ: {str(e)[:300]}")
    log.info("Run window finished, exiting (the workflow will restart).")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        sys.exit(0)
